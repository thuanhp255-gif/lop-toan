# -*- coding: utf-8 -*-
"""
du_lieu.py – Phần "bộ não" của ứng dụng: đọc / ghi file Excel và tính toán.
File này KHÔNG chứa giao diện. Giao diện nằm trong app.py.

Dữ liệu được lưu trong MỘT bảng tính – file Excel QuanLyLopToan.xlsx (khi chạy trên
máy) hoặc một Google Sheet (khi chạy online) – gồm các trang tính (sheet):
    HocSinh   – danh sách học sinh và thông tin cơ bản
    DiemDanh  – lịch sử điểm danh theo ngày
    Diem      – điểm các bài kiểm tra
    HocPhi    – học phí từng tháng của từng học sinh
    ChiPhi    – các khoản chi (để tính Thu – Chi)
    NhanXet   – nhận xét cuối tháng gửi phụ huynh
    CaiDat    – tên trung tâm, tên giáo viên, học phí mặc định...
Bạn có thể mở file này bằng Excel để xem / sửa tay (nhớ ĐÓNG Excel trước khi
bấm Lưu trong ứng dụng).
"""
from __future__ import annotations

import io
import re
import secrets
import shutil
import unicodedata
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ----------------------------------------------------------------------------
# 1. CẤU HÌNH CƠ BẢN
# ----------------------------------------------------------------------------
THU_MUC = Path(__file__).resolve().parent          # thư mục chứa ứng dụng
FILE_DU_LIEU = THU_MUC / "QuanLyLopToan.xlsx"      # file Excel dữ liệu
THU_MUC_SAO_LUU = THU_MUC / "sao_luu"              # nơi giữ bản sao lưu tự động

# Tên các cột trong từng sheet (đừng đổi tên cột trong Excel nhé)
COT = {
    "HocSinh":  ["STT", "Họ và tên", "Ngày sinh", "SĐT phụ huynh", "Nhóm/Lớp", "Ghi chú", "Tên đăng nhập", "Mật khẩu"],
    "DiemDanh": ["Ngày", "STT", "Họ và tên", "Trạng thái", "Ghi chú"],
    "Diem":     ["Ngày", "STT", "Họ và tên", "Loại bài", "Tên bài", "Điểm", "Nhận xét"],
    "HocPhi":   ["Tháng", "STT", "Họ và tên", "Phải đóng", "Đã đóng", "Ngày đóng", "Trạng thái", "Ghi chú"],
    "ChiPhi":   ["Ngày", "Nội dung", "Số tiền", "Ghi chú"],
    "NhanXet":  ["Tháng", "STT", "Họ và tên", "Nhận xét"],
    "CaiDat":   ["Mục", "Giá trị"],
}
COT_NGAY = {"DiemDanh": ["Ngày"], "Diem": ["Ngày"], "HocPhi": ["Ngày đóng"], "ChiPhi": ["Ngày"]}
COT_SO = {"HocSinh": ["STT"], "DiemDanh": ["STT"], "Diem": ["STT", "Điểm"],
          "HocPhi": ["STT", "Phải đóng", "Đã đóng"], "ChiPhi": ["Số tiền"], "NhanXet": ["STT"]}

TRANG_THAI_DD = ["Đi học", "Muộn", "Vắng"]
LOAI_BAI = ["Định kỳ", "Chuyên đề", "15 phút", "Khác"]
CAI_DAT_MAC_DINH = {
    "Tên trung tâm": "Trung tâm Toán",
    "Giáo viên": "Giáo viên Toán",
    "Môn": "Toán",
    "Học phí mặc định (đồng/tháng)": "600000",
    "Học sinh xem được học phí": "Có",
}


# ----------------------------------------------------------------------------
# 2. TIỆN ÍCH NHỎ
# ----------------------------------------------------------------------------
def thang_hien_tai() -> str:
    """Trả về tháng hiện tại dạng '2026-10'."""
    return date.today().strftime("%Y-%m")


def nhan_thang(thang: str) -> str:
    """'2026-10' -> 'Tháng 10/2026'."""
    try:
        y, m = str(thang).split("-")[:2]
        return f"Tháng {int(m)}/{y}"
    except Exception:
        return str(thang)


def dinh_dang_tien(so) -> str:
    """600000 -> '600.000 đ'."""
    try:
        return f"{int(round(float(so))):,}".replace(",", ".") + " đ"
    except Exception:
        return "0 đ"


def dinh_dang_diem(d) -> str:
    """8.5 -> '8,5' (kiểu Việt Nam)."""
    if d is None or (isinstance(d, float) and pd.isna(d)):
        return "—"
    s = f"{float(d):.2f}".rstrip("0").rstrip(".")
    return s.replace(".", ",")


def _ve_ngay(x):
    """Chuyển mọi kiểu ngày (Excel, chữ '05/10/2026', datetime) về date."""
    try:
        if x is None or x == "" or pd.isna(x):      # ô trống, NaN, NaT
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    try:
        return pd.to_datetime(str(x), dayfirst=True).date()
    except Exception:
        return None


def _ve_thang(x) -> str:
    """Chuẩn hoá tháng về dạng '2026-10' (chấp nhận cả '10/2026' hoặc ngày)."""
    if isinstance(x, (datetime, date)):
        return x.strftime("%Y-%m")
    s = str(x or "").strip()
    if "/" in s:                      # 10/2026
        p = s.split("/")
        if len(p) == 2:
            return f"{p[1]}-{int(p[0]):02d}"
    return s[:7]


# ----------------------------------------------------------------------------
# 3. CHUYỂN ĐỔI GIỮA "BẢNG" (pandas) VÀ "CÁC DÒNG" (Excel / Google Sheets)
# ----------------------------------------------------------------------------
def _chu(v) -> str:
    """Giá trị ô -> chữ gọn (7.0 -> '7', None -> '')."""
    if v is None:
        return ""
    try:
        if pd.isna(v):
            return ""
    except (TypeError, ValueError):
        pass
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def bang_tu_hang(ten: str, rows: list) -> pd.DataFrame:
    """Danh sách dòng (dòng đầu là tiêu đề) -> bảng pandas đủ cột, đúng kiểu."""
    cot = COT[ten]
    if not rows:
        return pd.DataFrame(columns=cot)
    tieu_de = [str(c).strip() if c is not None else "" for c in rows[0]]
    n = len(tieu_de)
    du_lieu = []
    for r in rows[1:]:
        r = list(r)[:n] + [None] * (n - len(r))          # Google Sheets cắt bớt ô trống cuối dòng
        if any(c not in (None, "") for c in r):
            du_lieu.append(r)
    df = pd.DataFrame(du_lieu, columns=tieu_de) if du_lieu else pd.DataFrame(columns=tieu_de)
    df = df.loc[:, ~df.columns.duplicated()]
    for c in cot:                                          # thêm cột còn thiếu
        if c not in df.columns:
            df[c] = None
    df = df[cot].copy()
    for c in COT_NGAY.get(ten, []):
        df[c] = df[c].map(_ve_ngay)
    for c in COT_SO.get(ten, []):
        df[c] = pd.to_numeric(df[c].map(_chu).str.replace(",", ".", regex=False).replace("", None), errors="coerce")
    if ten in ("HocPhi", "NhanXet"):
        df["Tháng"] = df["Tháng"].map(_ve_thang)
    if "STT" in df.columns:
        df = df[df["STT"].notna()].copy()
        df["STT"] = df["STT"].astype(int)
    for c in [c for c in cot if c not in COT_NGAY.get(ten, []) + COT_SO.get(ten, [])]:
        df[c] = df[c].map(_chu)
    return df.reset_index(drop=True)


def hang_tu_bang(ten: str, df: pd.DataFrame, kieu: str = "excel") -> list:
    """Bảng pandas -> danh sách dòng để ghi. kieu='chu': ngày thành chữ dd/mm/yyyy (cho Google Sheets)."""
    cot = COT[ten]
    hang = [list(cot)]
    if df is None:
        return hang
    for _, r in df.iterrows():
        dong = []
        for c in cot:
            v = r.get(c)
            try:
                trong = v is None or bool(pd.isna(v))
            except (TypeError, ValueError):
                trong = False
            if trong:
                v = None
            elif hasattr(v, "item"):                       # số kiểu numpy -> số Python
                v = v.item()
            if kieu == "chu":
                if v is None:
                    v = ""
                elif isinstance(v, (datetime, date)):
                    v = v.strftime("%d/%m/%Y")
                elif isinstance(v, float) and v.is_integer():
                    v = int(v)
            dong.append(v)
        hang.append(dong)
    return hang


# ----------------------------------------------------------------------------
# 4. NƠI LƯU DỮ LIỆU ("KHO"): Excel trên máy hoặc Google Sheets (online)
# ----------------------------------------------------------------------------
class FileDangMoLoi(Exception):
    """Báo lỗi khi file Excel đang mở trong Excel nên không ghi được."""


class KhoExcel:
    """Lưu trong file QuanLyLopToan.xlsx cạnh ứng dụng (khi chạy trên máy)."""
    ten = "Excel trên máy"

    def __init__(self, duong_dan: Path = FILE_DU_LIEU):
        self.duong_dan = Path(duong_dan)

    def ton_tai(self) -> bool:
        return self.duong_dan.exists()

    def moc(self) -> float:
        return self.duong_dan.stat().st_mtime

    def doc_hang(self) -> dict:
        wb = load_workbook(self.duong_dan, data_only=True)
        out = {ten: [list(r) for r in wb[ten].iter_rows(values_only=True)] if ten in wb.sheetnames else [] for ten in COT}
        wb.close()
        return out

    def _sao_luu(self):
        """Mỗi ngày giữ 1 bản sao lưu trong thư mục sao_luu/ (giữ 30 bản gần nhất)."""
        if not self.duong_dan.exists():
            return
        THU_MUC_SAO_LUU.mkdir(exist_ok=True)
        dich = THU_MUC_SAO_LUU / f"{self.duong_dan.stem}_{date.today():%Y-%m-%d}.xlsx"
        if not dich.exists():
            shutil.copy2(self.duong_dan, dich)
            for f in sorted(THU_MUC_SAO_LUU.glob(f"{self.duong_dan.stem}_*.xlsx"))[:-30]:
                f.unlink(missing_ok=True)

    def ghi(self, data: dict, cac_sheet=None):
        """Excel luôn ghi lại cả file (nhanh vì file nhỏ)."""
        self._sao_luu()
        wb = Workbook()
        wb.remove(wb.active)
        dau = PatternFill("solid", fgColor="1F4E79")
        for ten, cot in COT.items():
            ws = wb.create_sheet(ten)
            for dong in hang_tu_bang(ten, data.get(ten)):
                ws.append(dong)
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = dau
                cell.alignment = Alignment(vertical="center")
            ws.freeze_panes = "A2"
            for i, c in enumerate(cot, start=1):
                ws.column_dimensions[get_column_letter(i)].width = {"Họ và tên": 26, "Nhận xét": 50, "Ghi chú": 28,
                                                                     "Nội dung": 30, "Tên bài": 24}.get(c, 14)
                dinh_dang = "DD/MM/YYYY" if c in COT_NGAY.get(ten, []) else "#,##0" if c in ("Phải đóng", "Đã đóng", "Số tiền") else None
                if c == "Mật khẩu":
                    dinh_dang = "@"
                if dinh_dang:
                    for o in ws.iter_cols(min_col=i, max_col=i, min_row=2):
                        for x in o:
                            x.number_format = dinh_dang
        tam = self.duong_dan.with_suffix(".tmp.xlsx")
        try:
            wb.save(tam)
            tam.replace(self.duong_dan)
        except PermissionError as e:
            try:
                tam.unlink(missing_ok=True)
            except Exception:
                pass
            raise FileDangMoLoi("File Excel đang được mở bằng Excel. Hãy đóng Excel rồi thử lại.") from e


class KhoGoogle:
    """Lưu trong một Google Sheet (khi chạy online). bang_tinh: đối tượng gspread.Spreadsheet."""
    ten = "Google Sheets"

    def __init__(self, bang_tinh):
        self.sh = bang_tinh
        co_san = {ws.title for ws in self.sh.worksheets()}
        for ten, cot in COT.items():                       # tự tạo trang tính còn thiếu
            if ten not in co_san:
                ws = self.sh.add_worksheet(title=ten, rows=1000, cols=max(10, len(cot)))
                ws.update(values=[cot], range_name="A1")

    def ton_tai(self) -> bool:
        return True

    def doc_hang(self) -> dict:
        kq = self.sh.values_batch_get([f"'{t}'" for t in COT])        # 1 lần gọi cho cả 7 trang
        vung = kq.get("valueRanges", [])
        return {ten: (vung[i].get("values", []) if i < len(vung) else []) for i, ten in enumerate(COT)}

    def ghi(self, data: dict, cac_sheet=None):
        """Chỉ ghi các trang có thay đổi. Ghi đè trước, rồi xoá phần dòng thừa phía dưới (an toàn hơn xoá trước)."""
        ds = [t for t in (cac_sheet or COT) if t in COT]
        moi = {t: hang_tu_bang(t, data.get(t), "chu") for t in ds}
        self.sh.values_batch_update(body={"valueInputOption": "RAW",
                                          "data": [{"range": f"'{t}'!A1", "values": moi[t]} for t in ds]})
        self.sh.values_batch_clear(body={"ranges": [f"'{t}'!A{len(moi[t]) + 1}:Z" for t in ds]})


KHO = KhoExcel()          # app.py sẽ đổi sang KhoGoogle khi chạy online


def doc_tat_ca() -> dict[str, pd.DataFrame]:
    """Đọc toàn bộ dữ liệu. Trả về {tên trang tính: bảng}."""
    hang = KHO.doc_hang()
    data = {ten: bang_tu_hang(ten, hang.get(ten, [])) for ten in COT}
    cd = data["CaiDat"]
    for k, v in CAI_DAT_MAC_DINH.items():                  # bổ sung cài đặt còn thiếu
        if k not in set(cd["Mục"]):
            cd = pd.concat([cd, pd.DataFrame([{"Mục": k, "Giá trị": v}])], ignore_index=True)
    data["CaiDat"] = cd
    return data


def ghi_tat_ca(data: dict[str, pd.DataFrame], cac_sheet=None):
    """Ghi dữ liệu. cac_sheet: danh sách trang cần ghi (None = tất cả)."""
    KHO.ghi(data, cac_sheet)


def cai_dat(data) -> dict:
    return dict(zip(data["CaiDat"]["Mục"], data["CaiDat"]["Giá trị"]))


# ----------------------------------------------------------------------------
# 5. TÀI KHOẢN HỌC SINH + TẠO DỮ LIỆU LẦN ĐẦU
# ----------------------------------------------------------------------------
def bo_dau(s: str) -> str:
    """'Bảo Anh' -> 'bao anh'."""
    s = unicodedata.normalize("NFD", str(s)).replace("đ", "d").replace("Đ", "D")
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn")
    return re.sub(r"[^a-z0-9 ]", "", s.lower()).strip()


def ten_dang_nhap_mac_dinh(ho_ten: str, stt: int) -> str:
    """Tên (chữ cuối của họ tên, bỏ dấu) + số thứ tự 2 chữ số. 'Đoàn Trần Bảo An', 1 -> 'an01'."""
    tu = bo_dau(ho_ten).split()
    return (tu[-1] if tu else "hs") + f"{int(stt):02d}"


def cap_tai_khoan(hs: pd.DataFrame, cap_lai_mat_khau: bool = False) -> pd.DataFrame:
    """Điền tên đăng nhập + mật khẩu 2 chữ số cho học sinh nào còn trống."""
    hs = hs.copy()
    for c in ("Tên đăng nhập", "Mật khẩu"):
        if c not in hs.columns:
            hs[c] = ""
        hs[c] = hs[c].map(_chu)
    for i, r in hs.iterrows():
        if not r["Tên đăng nhập"]:
            hs.at[i, "Tên đăng nhập"] = ten_dang_nhap_mac_dinh(r["Họ và tên"], r["STT"])
        if cap_lai_mat_khau or not r["Mật khẩu"]:
            hs.at[i, "Mật khẩu"] = str(secrets.choice(range(10, 100)))
    hs["Tên đăng nhập"] = hs["Tên đăng nhập"].str.lower().str.replace(" ", "", regex=False)
    return hs


def kiem_tra_dang_nhap(data, ten: str, mat_khau: str, tk_gv: str, mk_gv: str):
    """Trả về ('gv', None), ('hs', STT) hoặc None nếu sai."""
    ten = bo_dau(ten).replace(" ", "")
    mat_khau = str(mat_khau).strip()
    if tk_gv and ten == bo_dau(tk_gv).replace(" ", "") and mk_gv and mat_khau == str(mk_gv):
        return ("gv", None)
    hs = data["HocSinh"]
    r = hs[hs["Tên đăng nhập"].str.lower() == ten]
    if len(r) and mat_khau and mat_khau.zfill(2) == str(r["Mật khẩu"].iloc[0]).strip().zfill(2):
        return ("hs", int(r["STT"].iloc[0]))
    return None


def doc_danh_sach_goc(nguon) -> pd.DataFrame:
    """
    Đọc file Excel danh sách gốc của bạn (có cột 'Họ và tên', tiêu đề có thể nằm ở
    dòng bất kỳ). Lấy thêm các cột nếu có: Ngày sinh, SĐT phụ huynh, Nhóm/Lớp.
    """
    wb = load_workbook(nguon, data_only=True)
    ket_qua = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        for i, r in enumerate(rows):
            ten_cot = [str(c).strip().lower() if c else "" for c in r]
            if "họ và tên" in ten_cot or "họ tên" in ten_cot:
                idx_ten = ten_cot.index("họ và tên") if "họ và tên" in ten_cot else ten_cot.index("họ tên")

                def tim(*tu):
                    for j, c in enumerate(ten_cot):
                        if any(t in c for t in tu):
                            return j
                    return None
                j_ns, j_sdt, j_nhom = tim("ngày sinh"), tim("sđt", "điện thoại", "phụ huynh"), tim("nhóm", "lớp")
                for r2 in rows[i + 1:]:
                    ten = r2[idx_ten] if idx_ten < len(r2) else None
                    if not ten or not str(ten).strip():
                        continue
                    ns = r2[j_ns] if j_ns is not None else ""
                    if isinstance(ns, (datetime, date)):
                        ns = ns.strftime("%d/%m/%Y")
                    ket_qua.append({"Họ và tên": " ".join(str(ten).split()), "Ngày sinh": ns or "",
                                    "SĐT phụ huynh": _chu(r2[j_sdt]) if j_sdt is not None else "",
                                    "Nhóm/Lớp": _chu(r2[j_nhom]) if j_nhom is not None else ws.title, "Ghi chú": ""})
                break
        if ket_qua:
            break
    wb.close()
    if not ket_qua:
        raise ValueError("Không tìm thấy cột 'Họ và tên' trong file.")
    df = pd.DataFrame(ket_qua)
    df.insert(0, "STT", range(1, len(df) + 1))
    df = cap_tai_khoan(df)
    return df[COT["HocSinh"]]


def du_lieu_moi(ds_hoc_sinh: pd.DataFrame) -> dict:
    data = {ten: pd.DataFrame(columns=cot) for ten, cot in COT.items()}
    data["HocSinh"] = ds_hoc_sinh
    data["CaiDat"] = pd.DataFrame([{"Mục": k, "Giá trị": v} for k, v in CAI_DAT_MAC_DINH.items()])
    return data


def tao_file_moi(ds_hoc_sinh: pd.DataFrame):
    ghi_tat_ca(du_lieu_moi(ds_hoc_sinh))


# ----------------------------------------------------------------------------
# 6. NGHIỆP VỤ: ĐIỂM DANH
# ----------------------------------------------------------------------------
def ten_hs(data, stt) -> str:
    hs = data["HocSinh"]
    r = hs[hs["STT"] == int(stt)]
    return r["Họ và tên"].iloc[0] if len(r) else ""


def diem_danh_ngay(data, ngay: date) -> dict[int, str]:
    """{STT: trạng thái} của một ngày."""
    dd = data["DiemDanh"]
    dd = dd[dd["Ngày"] == ngay]
    return dict(zip(dd["STT"].astype(int), dd["Trạng thái"]))


def ghi_diem_danh(data, ngay: date, stt: int, trang_thai: str | None):
    """Cập nhật 1 ô điểm danh (trạng thái None = xoá)."""
    dd = data["DiemDanh"]
    dd = dd[~((dd["Ngày"] == ngay) & (dd["STT"] == int(stt)))]
    if trang_thai:
        dd = pd.concat([dd, pd.DataFrame([{"Ngày": ngay, "STT": int(stt), "Họ và tên": ten_hs(data, stt),
                                           "Trạng thái": trang_thai, "Ghi chú": ""}])], ignore_index=True)
    data["DiemDanh"] = dd.sort_values(["Ngày", "STT"]).reset_index(drop=True)


def thong_ke_chuyen_can(data, thang: str | None = None) -> pd.DataFrame:
    """Bảng số buổi Đi học / Muộn / Vắng của từng học sinh (lọc theo tháng nếu có)."""
    dd = data["DiemDanh"].copy()
    if thang:
        dd = dd[dd["Ngày"].map(lambda d: d.strftime("%Y-%m") if d else "") == thang]
    hs = data["HocSinh"][["STT", "Họ và tên"]].copy()
    for t in TRANG_THAI_DD:
        dem = dd[dd["Trạng thái"] == t].groupby("STT").size()
        hs[t] = hs["STT"].map(dem).fillna(0).astype(int)
    tong = hs[TRANG_THAI_DD].sum(axis=1)
    hs["Số buổi"] = tong
    hs["Chuyên cần (%)"] = ((hs["Đi học"] + hs["Muộn"]) / tong.where(tong > 0) * 100).round(0)
    return hs


# ----------------------------------------------------------------------------
# 7. NGHIỆP VỤ: ĐIỂM SỐ
# ----------------------------------------------------------------------------
def them_diem(data, ngay: date, stt: int, loai: str, ten_bai: str, diem: float, nhan_xet: str = ""):
    """Thêm (hoặc ghi đè nếu trùng ngày + tên bài) một điểm."""
    d = data["Diem"]
    d = d[~((d["Ngày"] == ngay) & (d["STT"] == int(stt)) & (d["Tên bài"] == ten_bai))]
    moi = {"Ngày": ngay, "STT": int(stt), "Họ và tên": ten_hs(data, stt), "Loại bài": loai,
           "Tên bài": ten_bai, "Điểm": float(diem), "Nhận xét": nhan_xet or ""}
    data["Diem"] = pd.concat([d, pd.DataFrame([moi])], ignore_index=True).sort_values(["Ngày", "STT"]).reset_index(drop=True)


def diem_cua_hs(data, stt: int, thang: str | None = None) -> pd.DataFrame:
    d = data["Diem"]
    d = d[(d["STT"] == int(stt)) & d["Điểm"].notna()].copy()
    if thang:
        d = d[d["Ngày"].map(lambda x: x.strftime("%Y-%m") if x else "") == thang]
    return d.sort_values("Ngày").reset_index(drop=True)


def tom_tat_diem(d: pd.DataFrame) -> dict:
    """Điểm TB chung, TB theo loại, bài gần nhất và mức thay đổi so với bài trước."""
    kq = {"so_bai": len(d), "tb": None, "tb_loai": {}, "gan_nhat": None, "thay_doi": None, "xu_huong": "Chưa đủ dữ liệu"}
    if d.empty:
        return kq
    kq["tb"] = round(float(d["Điểm"].mean()), 2)
    for loai, g in d.groupby("Loại bài"):
        kq["tb_loai"][loai] = round(float(g["Điểm"].mean()), 2)
    kq["gan_nhat"] = float(d["Điểm"].iloc[-1])
    if len(d) >= 2:
        kq["thay_doi"] = round(float(d["Điểm"].iloc[-1] - d["Điểm"].iloc[-2]), 2)
        # Xu hướng: so TB nửa sau với nửa đầu các bài
        n = len(d) // 2
        chenh = d["Điểm"].iloc[n:].mean() - d["Điểm"].iloc[:n].mean() if n else 0
        kq["xu_huong"] = "Tiến bộ" if chenh >= 0.5 else ("Đi xuống" if chenh <= -0.5 else "Ổn định")
    return kq


def xep_loai(tb) -> str:
    if tb is None:
        return "—"
    return "Giỏi" if tb >= 8 else "Khá" if tb >= 6.5 else "Trung bình" if tb >= 5 else "Yếu"


# ----------------------------------------------------------------------------
# 8. NGHIỆP VỤ: HỌC PHÍ & THU – CHI
# ----------------------------------------------------------------------------
def trang_thai_hp(phai, da) -> str:
    phai = 0 if pd.isna(phai) else float(phai)
    da = 0 if pd.isna(da) else float(da)
    if phai <= 0:
        return "Miễn phí"
    if da >= phai:
        return "Đã đóng đủ"
    if da > 0:
        return "Còn nợ"
    return "Chưa đóng"


def bang_hoc_phi(data, thang: str) -> pd.DataFrame:
    """Bảng học phí của tháng; học sinh chưa có dòng sẽ được tạo với học phí mặc định."""
    mac_dinh = float(str(cai_dat(data).get("Học phí mặc định (đồng/tháng)", 0)).replace(".", "").replace(",", "") or 0)
    hp = data["HocPhi"]
    hp_t = hp[hp["Tháng"] == thang].set_index("STT")
    dong = []
    for _, h in data["HocSinh"].iterrows():
        stt = int(h["STT"])
        if stt in hp_t.index:
            r = hp_t.loc[stt]
            if isinstance(r, pd.DataFrame):
                r = r.iloc[-1]
            dong.append({"STT": stt, "Họ và tên": h["Họ và tên"], "Phải đóng": r["Phải đóng"] if pd.notna(r["Phải đóng"]) else mac_dinh,
                         "Đã đóng": r["Đã đóng"] if pd.notna(r["Đã đóng"]) else 0, "Ngày đóng": r["Ngày đóng"], "Ghi chú": r["Ghi chú"]})
        else:
            dong.append({"STT": stt, "Họ và tên": h["Họ và tên"], "Phải đóng": mac_dinh, "Đã đóng": 0, "Ngày đóng": None, "Ghi chú": ""})
    df = pd.DataFrame(dong)
    df["Ngày đóng"] = pd.Series([_ve_ngay(x) for x in df["Ngày đóng"]], dtype=object)
    df["Trạng thái"] = [trang_thai_hp(p, d) for p, d in zip(df["Phải đóng"], df["Đã đóng"])]
    df["Còn nợ"] = (df["Phải đóng"].fillna(0) - df["Đã đóng"].fillna(0)).clip(lower=0)
    return df


def luu_hoc_phi(data, thang: str, bang: pd.DataFrame):
    hp = data["HocPhi"]
    hp = hp[hp["Tháng"] != thang]
    moi = bang.copy()
    moi["Tháng"] = thang
    moi["Trạng thái"] = [trang_thai_hp(p, d) for p, d in zip(moi["Phải đóng"], moi["Đã đóng"])]
    moi["Ngày đóng"] = moi["Ngày đóng"].map(_ve_ngay)
    moi["Ghi chú"] = moi["Ghi chú"].fillna("").astype(str)
    data["HocPhi"] = pd.concat([hp, moi[COT["HocPhi"]]], ignore_index=True).sort_values(["Tháng", "STT"]).reset_index(drop=True)


def chi_trong_thang(data, thang: str) -> pd.DataFrame:
    c = data["ChiPhi"]
    return c[c["Ngày"].map(lambda d: d.strftime("%Y-%m") if d else "") == thang].reset_index(drop=True)


# ----------------------------------------------------------------------------
# 9. NHẬN XÉT & BÁO CÁO PHỤ HUYNH
# ----------------------------------------------------------------------------
def lay_nhan_xet(data, thang: str, stt: int) -> str:
    nx = data["NhanXet"]
    r = nx[(nx["Tháng"] == thang) & (nx["STT"] == int(stt))]
    return r["Nhận xét"].iloc[-1] if len(r) else ""


def ghi_nhan_xet(data, thang: str, stt: int, noi_dung: str):
    nx = data["NhanXet"]
    nx = nx[~((nx["Tháng"] == thang) & (nx["STT"] == int(stt)))]
    if noi_dung.strip():
        nx = pd.concat([nx, pd.DataFrame([{"Tháng": thang, "STT": int(stt), "Họ và tên": ten_hs(data, stt),
                                           "Nhận xét": noi_dung.strip()}])], ignore_index=True)
    data["NhanXet"] = nx.sort_values(["Tháng", "STT"]).reset_index(drop=True)


def thang_truoc(thang: str) -> str:
    y, m = map(int, thang.split("-"))
    return f"{y - 1}-12" if m == 1 else f"{y}-{m - 1:02d}"


def noi_dung_bao_cao(data, stt: int, thang: str, kem_hoc_phi: bool = True) -> dict:
    """Gom mọi số liệu cần cho báo cáo 1 học sinh trong 1 tháng."""
    cd = cai_dat(data)
    cc = thong_ke_chuyen_can(data, thang)
    cc = cc[cc["STT"] == int(stt)].iloc[0]
    d = diem_cua_hs(data, stt, thang)
    tt = tom_tat_diem(d)
    tb_truoc = tom_tat_diem(diem_cua_hs(data, stt, thang_truoc(thang)))["tb"]
    hp = None
    if kem_hoc_phi:
        b = bang_hoc_phi(data, thang)
        hp = b[b["STT"] == int(stt)].iloc[0].to_dict()
    hs = data["HocSinh"][data["HocSinh"]["STT"] == int(stt)].iloc[0]
    return {"trung_tam": cd.get("Tên trung tâm", ""), "giao_vien": cd.get("Giáo viên", ""), "mon": cd.get("Môn", "Toán"),
            "ten": hs["Họ và tên"], "nhom": hs["Nhóm/Lớp"], "thang": thang, "chuyen_can": cc.to_dict(),
            "diem": d, "tom_tat": tt, "tb_truoc": tb_truoc, "hoc_phi": hp, "nhan_xet": lay_nhan_xet(data, thang, stt)}


def bao_cao_van_ban(bc: dict) -> str:
    """Báo cáo dạng chữ – dán thẳng vào Zalo / tin nhắn."""
    cc, tt = bc["chuyen_can"], bc["tom_tat"]
    dong = [f"BÁO CÁO HỌC TẬP MÔN {bc['mon'].upper()} – {nhan_thang(bc['thang']).upper()}",
            f"{bc['trung_tam']}",
            "",
            f"Học sinh: {bc['ten']}" + (f" ({bc['nhom']})" if bc["nhom"] else ""),
            "",
            "1. Chuyên cần",
            f"   Đi học {int(cc['Đi học'])} · Muộn {int(cc['Muộn'])} · Vắng {int(cc['Vắng'])} / {int(cc['Số buổi'])} buổi"
            + (f" (chuyên cần {int(cc['Chuyên cần (%)'])}%)" if pd.notna(cc["Chuyên cần (%)"]) else ""),
            "",
            "2. Kết quả kiểm tra"]
    if tt["so_bai"]:
        for _, r in bc["diem"].iterrows():
            ngay = r["Ngày"].strftime("%d/%m") if r["Ngày"] else ""
            dong.append(f"   - {ngay} · {r['Tên bài'] or r['Loại bài']} ({r['Loại bài']}): {dinh_dang_diem(r['Điểm'])}"
                        + (f" – {r['Nhận xét']}" if r["Nhận xét"] else ""))
        so_sanh = ""
        if bc["tb_truoc"] is not None:
            ch = tt["tb"] - bc["tb_truoc"]
            so_sanh = f" ({'tăng' if ch > 0 else 'giảm' if ch < 0 else 'bằng'} {dinh_dang_diem(abs(ch)) if ch else ''} so với tháng trước)".replace("  ", " ")
        dong.append(f"   Điểm trung bình: {dinh_dang_diem(tt['tb'])} – {xep_loai(tt['tb'])}{so_sanh}")
    else:
        dong.append("   Tháng này chưa có bài kiểm tra.")
    dong += ["", "3. Nhận xét của giáo viên", f"   {bc['nhan_xet'] or '(chưa có nhận xét)'}"]
    if bc["hoc_phi"]:
        hp = bc["hoc_phi"]
        dong += ["", "4. Học phí", f"   {hp['Trạng thái']}" + (f" – còn {dinh_dang_tien(hp['Còn nợ'])}" if hp["Còn nợ"] > 0 else "")]
    dong += ["", f"Trân trọng,", f"{bc['giao_vien']}"]
    return "\n".join(dong)


def bao_cao_anh(bc: dict) -> bytes:
    """Báo cáo dạng ảnh PNG (gửi Zalo đẹp hơn). Dùng font DejaVu có sẵn hỗ trợ tiếng Việt."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    XANH, CHU, MO, DO = "#1F4E79", "#1B2430", "#6B7785", "#C0392B"
    plt.rcParams["font.family"] = "DejaVu Sans"
    fig = plt.figure(figsize=(8.27, 11.0), dpi=150)
    fig.patch.set_facecolor("white")

    def chu(x, y, s, **kw):
        fig.text(x, y, s, transform=fig.transFigure, **kw)

    # Đầu trang
    fig.add_artist(plt.Rectangle((0, 0.9), 1, 0.1, color=XANH, transform=fig.transFigure))
    chu(0.06, 0.955, f"BÁO CÁO HỌC TẬP MÔN {bc['mon'].upper()}", color="white", fontsize=17, weight="bold", va="center")
    chu(0.06, 0.92, f"{bc['trung_tam']} · {nhan_thang(bc['thang'])}", color="white", fontsize=10.5, va="center")
    chu(0.06, 0.865, bc["ten"], color=CHU, fontsize=19, weight="bold", va="center")
    if bc["nhom"]:
        chu(0.06, 0.838, f"Nhóm/Lớp: {bc['nhom']}", color=MO, fontsize=10, va="center")

    # 3 ô số liệu
    cc, tt = bc["chuyen_can"], bc["tom_tat"]
    o = [("Chuyên cần", f"{int(cc['Đi học']) + int(cc['Muộn'])}/{int(cc['Số buổi'])} buổi",
          f"Vắng {int(cc['Vắng'])} · Muộn {int(cc['Muộn'])}"),
         ("Điểm trung bình", dinh_dang_diem(tt["tb"]), xep_loai(tt["tb"]) if tt["tb"] is not None else "Chưa có bài"),
         ("Xu hướng", tt["xu_huong"], f"{tt['so_bai']} bài kiểm tra")]
    for i, (nhan, so, phu) in enumerate(o):
        x = 0.06 + i * 0.30
        fig.add_artist(plt.Rectangle((x, 0.70), 0.27, 0.11, fill=False, ec="#D5DCE4", lw=1, transform=fig.transFigure))
        chu(x + 0.015, 0.79, nhan, fontsize=9, color=MO)
        chu(x + 0.015, 0.745, so, fontsize=16 if len(so) < 12 else 12, weight="bold", color=DO if i == 1 else CHU)
        chu(x + 0.015, 0.712, phu, fontsize=8.5, color=MO)

    # Biểu đồ điểm
    ax = fig.add_axes([0.09, 0.47, 0.85, 0.19])
    d = bc["diem"]
    if len(d):
        x = list(range(len(d)))
        ax.plot(x, d["Điểm"], color=DO, marker="o", lw=2)
        for xi, yi in zip(x, d["Điểm"]):
            ax.annotate(dinh_dang_diem(yi), (xi, yi), textcoords="offset points", xytext=(0, 7), ha="center", fontsize=8, color=DO)
        ax.set_xticks(x)
        ax.set_xticklabels([r["Ngày"].strftime("%d/%m") if r["Ngày"] else "" for _, r in d.iterrows()], fontsize=8)
        if tt["tb"] is not None:
            ax.axhline(tt["tb"], color=XANH, ls="--", lw=1)
            ax.text(len(d) - 0.5, tt["tb"] + 0.2, f"TB {dinh_dang_diem(tt['tb'])}", color=XANH, fontsize=8, ha="right")
    else:
        ax.text(0.5, 0.5, "Chưa có bài kiểm tra trong tháng", ha="center", va="center", transform=ax.transAxes, color=MO)
        ax.set_xticks([])
    ax.set_ylim(0, 10.8)
    ax.set_yticks([0, 5, 10])
    ax.grid(axis="y", color="#EEF1F5")
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.set_title("Điểm các bài kiểm tra (thang 10)", loc="left", fontsize=10, color=CHU)

    # Danh sách bài + nhận xét
    y = 0.42
    chu(0.06, y, "Chi tiết bài kiểm tra", fontsize=11, weight="bold", color=CHU)
    y -= 0.028
    for _, r in d.tail(8).iterrows():
        ngay = r["Ngày"].strftime("%d/%m") if r["Ngày"] else ""
        chu(0.06, y, f"{ngay}  {(r['Tên bài'] or r['Loại bài'])[:40]} ({r['Loại bài']})", fontsize=9, color=CHU)
        chu(0.94, y, dinh_dang_diem(r["Điểm"]), fontsize=10, weight="bold", color=DO, ha="right")
        y -= 0.024
    y -= 0.015
    chu(0.06, y, "Nhận xét của giáo viên", fontsize=11, weight="bold", color=CHU)
    import textwrap
    for dong in textwrap.wrap(bc["nhan_xet"] or "(chưa có nhận xét)", 95)[:7]:
        y -= 0.024
        chu(0.06, y, dong, fontsize=9.5, color=CHU)
    if bc["hoc_phi"]:
        hp = bc["hoc_phi"]
        y -= 0.04
        chu(0.06, y, "Học phí: " + hp["Trạng thái"] + (f" – còn {dinh_dang_tien(hp['Còn nợ'])}" if hp["Còn nợ"] > 0 else ""),
            fontsize=10, color=DO if hp["Còn nợ"] > 0 else "#1E8449", weight="bold")
    chu(0.94, 0.04, f"Giáo viên: {bc['giao_vien']}", fontsize=9.5, color=MO, ha="right")
    chu(0.06, 0.04, f"Ngày lập: {date.today():%d/%m/%Y}", fontsize=9, color=MO)

    buf = io.BytesIO()
    fig.savefig(buf, format="png", facecolor="white")
    plt.close(fig)
    return buf.getvalue()
