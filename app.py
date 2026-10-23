# -*- coding: utf-8 -*-
"""
app.py – Giao diện ứng dụng QUẢN LÝ LỚP TOÁN (Streamlit).
Chạy được 2 cách: trên máy (lưu Excel) hoặc ONLINE trên Streamlit Community Cloud (lưu Google Sheets,
mỗi học sinh 1 tài khoản: tên + số thứ tự, mật khẩu 2 chữ số).

Cách chạy:   python -m streamlit run app.py
(hoặc bấm đúp file  chay_app.bat  trên Windows /  chay_app.command  trên Mac)

Ứng dụng có 4 màn hình chính (chọn ở thanh bên trái):
    1. Tổng quan & Điểm danh
    2. Nhập & theo dõi điểm
    3. Học phí & Thu – Chi
    4. Báo cáo phụ huynh
và trang Cài đặt (tên trung tâm, tên giáo viên, học phí mặc định).
Học sinh đăng nhập chỉ thấy điểm, chuyên cần, nhận xét (và học phí nếu bật) của chính mình.
"""
import io
import time
import zipfile
from datetime import date

import altair as alt
import pandas as pd
import streamlit as st

import du_lieu as dl

# ============================================================================
# CÀI ĐẶT TRANG + GIAO DIỆN TỐI GIẢN
# ============================================================================
st.set_page_config(page_title="Quản lý lớp Toán", page_icon="📐", layout="wide")

# Tương thích nhiều phiên bản Streamlit: bản mới dùng width="stretch", bản cũ dùng use_container_width
try:
    _PB = tuple(int(x) for x in st.__version__.split(".")[:2])
except Exception:
    _PB = (1, 50)
FULL = {"width": "stretch"} if _PB >= (1, 49) else {"use_container_width": True}
st.markdown("""
<style>
  .block-container {padding-top: 1.6rem; max-width: 1200px;}
  h1 {font-size: 1.7rem !important; margin-bottom: .2rem;}
  h3 {font-size: 1.15rem !important;}
  div[data-testid="stMetricValue"] {font-size: 1.7rem;}
  .ten-hs {padding-top: .45rem; font-weight: 500;}
  .stt {padding-top: .45rem; color: #6B7785;}
  .nho {color: #6B7785; font-size: .9rem;}
  div[role="radiogroup"] {gap: .4rem;}
</style>
""", unsafe_allow_html=True)


# ============================================================================
# CHỌN NƠI LƯU DỮ LIỆU
#   • Có file .streamlit/secrets.toml (hoặc Secrets trên Streamlit Cloud) chứa [gsheets]
#     → chạy ONLINE, lưu vào Google Sheets, bắt buộc đăng nhập.
#   • Không có → chạy trên máy, lưu vào QuanLyLopToan.xlsx, vào thẳng trang giáo viên.
# ============================================================================
def _doc_bi_mat() -> dict:
    try:
        return {k: st.secrets[k] for k in st.secrets.keys()}
    except Exception:
        return {}


BI_MAT = _doc_bi_mat()
ONLINE = "gsheets" in BI_MAT


@st.cache_resource(show_spinner="Đang kết nối Google Sheets…")
def _ket_noi_google():
    """Kết nối Google Sheets bằng tài khoản dịch vụ (service account) trong Secrets."""
    import json
    import gspread
    if "gcp_json" in BI_MAT:                       # cách dễ: dán nguyên nội dung file JSON
        thong_tin = json.loads(BI_MAT["gcp_json"])
    else:
        thong_tin = dict(BI_MAT["gcp_service_account"])
    gc = gspread.service_account_from_dict(thong_tin)
    dia_chi = str(BI_MAT["gsheets"]["url"]).strip()
    bang = gc.open_by_url(dia_chi) if dia_chi.startswith("http") else gc.open_by_key(dia_chi)
    return dl.KhoGoogle(bang)


if ONLINE:
    try:
        dl.KHO = _ket_noi_google()
    except Exception as e:
        st.error("Không kết nối được Google Sheets. Kiểm tra lại mục Secrets (đường link bảng tính, "
                 "khoá JSON) và đã chia sẻ bảng tính cho email của tài khoản dịch vụ với quyền Người chỉnh sửa.")
        st.caption(f"Chi tiết lỗi: {type(e).__name__}: {e}")
        st.stop()
    GV = BI_MAT.get("giao_vien", {})
    TK_GV, MK_GV = str(GV.get("tai_khoan", "giaovien")), str(GV.get("mat_khau", ""))
    if not MK_GV:
        st.error("Chưa đặt mật khẩu giáo viên trong Secrets (mục [giao_vien] → mat_khau).")
        st.stop()


# ============================================================================
# ĐỌC / GHI DỮ LIỆU
# ============================================================================
@st.cache_data(show_spinner=False, ttl=60)
def _doc_du_lieu(moc: float):
    """Đọc dữ liệu. Excel: đọc lại khi file đổi. Google: đọc lại sau 60 giây hoặc ngay sau khi lưu."""
    return dl.doc_tat_ca()


def nap_du_lieu() -> dict:
    return _doc_du_lieu(0.0 if ONLINE else dl.KHO.moc())


def luu(data: dict, thong_bao: str, cac_sheet=None) -> bool:
    """Ghi dữ liệu (chỉ các trang tính cac_sheet nếu có) và báo kết quả."""
    try:
        dl.ghi_tat_ca(data, cac_sheet)
    except dl.FileDangMoLoi as e:
        st.error(f"⚠️ {e}")
        return False
    except Exception as e:                           # lỗi mạng / Google
        st.error(f"⚠️ Chưa lưu được: {type(e).__name__}: {e}. Kiểm tra mạng rồi bấm lưu lại.")
        return False
    _doc_du_lieu.clear()
    st.toast(f"✅ {thong_bao}")
    return True


def nhan_hs(data, stt) -> str:
    return f"{int(stt)}. {dl.ten_hs(data, stt)}"


# ============================================================================
# THIẾT LẬP LẦN ĐẦU: NHẬP DANH SÁCH HỌC SINH
# ============================================================================
def _doc_file_du_lieu_day_du(nguon):
    """Nếu file tải lên là file dữ liệu của ứng dụng (có trang HocSinh…) thì lấy nguyên mọi trang."""
    from openpyxl import load_workbook
    wb = load_workbook(nguon, data_only=True)
    if "HocSinh" not in wb.sheetnames:
        return None
    hang = {ten: [list(r) for r in wb[ten].iter_rows(values_only=True)] if ten in wb.sheetnames else [] for ten in dl.COT}
    d = {ten: dl.bang_tu_hang(ten, hang[ten]) for ten in dl.COT}
    d["HocSinh"] = dl.cap_tai_khoan(d["HocSinh"])
    if d["CaiDat"].empty:
        d["CaiDat"] = pd.DataFrame([{"Mục": k, "Giá trị": v} for k, v in dl.CAI_DAT_MAC_DINH.items()])
    return d


def man_hinh_khoi_tao():
    st.title("📐 Quản lý lớp Toán – Thiết lập lần đầu")
    st.write(f"Nơi lưu dữ liệu (**{dl.KHO.ten}**) chưa có danh sách học sinh. Hãy chọn file Excel danh sách "
             "(hoặc file dữ liệu QuanLyLopToan.xlsx) để ứng dụng tạo dữ liệu.")
    nguon = None
    if not ONLINE:
        ung_vien = [f for f in dl.THU_MUC.glob("*.xlsx") if f.name != dl.FILE_DU_LIEU.name and not f.name.startswith("~$")]
        if ung_vien:
            nguon = dl.THU_MUC / st.selectbox("File Excel tìm thấy trong thư mục ứng dụng", [f.name for f in ung_vien])
    tai_len = st.file_uploader("Chọn file Excel từ máy", type=["xlsx"])
    if tai_len is not None:
        nguon = io.BytesIO(tai_len.getvalue())
    if nguon is None:
        return
    try:
        day_du = _doc_file_du_lieu_day_du(nguon)
        if hasattr(nguon, "seek"):
            nguon.seek(0)
        data_moi = day_du if day_du is not None else dl.du_lieu_moi(dl.doc_danh_sach_goc(nguon))
    except Exception as e:
        st.error(f"Không đọc được file: {e}")
        return
    ds = data_moi["HocSinh"]
    st.success(f"Tìm thấy {len(ds)} học sinh" + (" cùng toàn bộ dữ liệu điểm, điểm danh, học phí." if day_du is not None else "."))
    st.dataframe(ds, hide_index=True, **FULL, height=300)
    if st.button("Tạo dữ liệu và bắt đầu", type="primary"):
        if luu(data_moi, "Đã tạo dữ liệu"):
            st.rerun()


if not ONLINE and not dl.KHO.ton_tai():
    man_hinh_khoi_tao()
    st.stop()

data = nap_du_lieu()
cd = dl.cai_dat(data)
hoc_sinh = data["HocSinh"].sort_values("STT").reset_index(drop=True)
DS_STT = hoc_sinh["STT"].astype(int).tolist()


def chon_thang(nhan: str, key: str) -> str:
    """Ô chọn tháng: 12 tháng gần nhất + các tháng đã có dữ liệu."""
    hom_nay = date.today()
    ds = set()
    for i in range(12):
        y, m = hom_nay.year, hom_nay.month - i
        while m <= 0:
            m += 12
            y -= 1
        ds.add(f"{y}-{m:02d}")
    for ten, cot in (("HocPhi", "Tháng"), ("NhanXet", "Tháng")):
        ds |= set(data[ten][cot].dropna().astype(str))
    for ten in ("DiemDanh", "Diem", "ChiPhi"):
        ds |= {d.strftime("%Y-%m") for d in data[ten]["Ngày"].dropna()}
    ds = sorted({x for x in ds if len(x) == 7}, reverse=True)
    return st.selectbox(nhan, ds, index=ds.index(dl.thang_hien_tai()) if dl.thang_hien_tai() in ds else 0,
                        format_func=dl.nhan_thang, key=key)


# ============================================================================
# MÀN HÌNH 1 – TỔNG QUAN & ĐIỂM DANH
# ============================================================================
def _khoa_dd(ngay, stt):
    return f"dd_{ngay.isoformat()}_{stt}"


def _cb_luu_diem_danh(ngay, ds_stt):
    """Lưu một lần mọi lựa chọn điểm danh trên màn hình (tiết kiệm lượt ghi Google Sheets)."""
    d = nap_du_lieu()
    for stt in ds_stt:
        dl.ghi_diem_danh(d, ngay, stt, st.session_state.get(_khoa_dd(ngay, stt)))
    luu(d, f"Đã lưu điểm danh ngày {ngay:%d/%m/%Y}", ["DiemDanh"])


def _cb_ca_lop(ngay, ds_stt):
    d = nap_du_lieu()
    hien_tai = dl.diem_danh_ngay(d, ngay)
    for stt in ds_stt:
        if stt not in hien_tai:
            dl.ghi_diem_danh(d, ngay, stt, "Đi học")
            st.session_state[_khoa_dd(ngay, stt)] = "Đi học"
    luu(d, "Đã đánh dấu cả lớp đi học", ["DiemDanh"])


def _cb_xoa_ngay(ngay, ds_stt):
    d = nap_du_lieu()
    for stt in ds_stt:
        dl.ghi_diem_danh(d, ngay, stt, None)
        st.session_state[_khoa_dd(ngay, stt)] = None
    luu(d, f"Đã xoá điểm danh ngày {ngay:%d/%m/%Y}", ["DiemDanh"])


def man_hinh_diem_danh():
    st.title("Tổng quan & Điểm danh")
    tab_dd, tab_ls, tab_hs = st.tabs(["✅ Điểm danh hôm nay", "📅 Lịch sử chuyên cần", "👥 Thông tin học sinh"])

    # ---------------- Điểm danh ----------------
    with tab_dd:
        c1, c2, c3 = st.columns([1.2, 1.2, 3])
        ngay = c1.date_input("Ngày học", date.today(), format="DD/MM/YYYY", key="ngay_dd")
        nhom_ds = sorted({n for n in hoc_sinh["Nhóm/Lớp"] if n})
        nhom = c2.selectbox("Nhóm/Lớp", ["Tất cả"] + nhom_ds) if len(nhom_ds) > 1 else "Tất cả"
        ds = hoc_sinh if nhom == "Tất cả" else hoc_sinh[hoc_sinh["Nhóm/Lớp"] == nhom]
        ds_stt = ds["STT"].astype(int).tolist()
        dd = dl.diem_danh_ngay(data, ngay)

        di = sum(1 for s in ds_stt if dd.get(s) == "Đi học")
        muon = sum(1 for s in ds_stt if dd.get(s) == "Muộn")
        vang = sum(1 for s in ds_stt if dd.get(s) == "Vắng")
        k = st.columns(5)
        k[0].metric("Sĩ số", len(ds_stt))
        k[1].metric("Đi học", di)
        k[2].metric("Muộn", muon)
        k[3].metric("Vắng", vang)
        k[4].metric("Chưa điểm danh", len(ds_stt) - di - muon - vang)

        b1, b2, _ = st.columns([1.4, 1.4, 4])
        b1.button("✔ Cả lớp đi học", on_click=_cb_ca_lop, args=(ngay, ds_stt), **FULL,
                  help="Đánh dấu 'Đi học' cho các em chưa được điểm danh")
        b2.button("Xoá điểm danh ngày này", on_click=_cb_xoa_ngay, args=(ngay, ds_stt), **FULL)
        cho_luu = [s for s in ds_stt if st.session_state.get(_khoa_dd(ngay, s), dd.get(s)) != dd.get(s)]
        st.caption("Chọn **Đi học / Muộn / Vắng** cho từng em rồi bấm **Lưu điểm danh** (nút ở cuối danh sách).")
        if cho_luu:
            st.warning(f"Có {len(cho_luu)} thay đổi chưa lưu.")

        # Chia 2 cột cho gọn màn hình
        nua = (len(ds) + 1) // 2
        cot_trai, cot_phai = st.columns(2, gap="large")
        for i, (_, hs) in enumerate(ds.iterrows()):
            stt = int(hs["STT"])
            khoa = _khoa_dd(ngay, stt)
            if khoa not in st.session_state:
                st.session_state[khoa] = dd.get(stt)
            with (cot_trai if i < nua else cot_phai):
                a, b, c = st.columns([0.35, 2.2, 2.6])
                a.markdown(f"<div class='stt'>{stt}</div>", unsafe_allow_html=True)
                b.markdown(f"<div class='ten-hs'>{hs['Họ và tên']}</div>", unsafe_allow_html=True)
                c.radio(f"Điểm danh {hs['Họ và tên']}", dl.TRANG_THAI_DD, key=khoa, horizontal=True,
                        label_visibility="collapsed")
        st.button(f"💾 Lưu điểm danh{f' ({len(cho_luu)} thay đổi)' if cho_luu else ''}", type="primary",
                  on_click=_cb_luu_diem_danh, args=(ngay, ds_stt), key="luu_dd")

    # ---------------- Lịch sử ----------------
    with tab_ls:
        thang = chon_thang("Tháng", "thang_ls")
        tk = dl.thong_ke_chuyen_can(data, thang)
        dd_t = data["DiemDanh"][data["DiemDanh"]["Ngày"].map(lambda d: d.strftime("%Y-%m") if d else "") == thang]
        so_ngay = dd_t["Ngày"].nunique()
        k = st.columns(4)
        k[0].metric("Số buổi đã điểm danh", so_ngay)
        tong = tk["Số buổi"].sum()
        k[1].metric("Chuyên cần cả lớp", f"{round((tk['Đi học'].sum() + tk['Muộn'].sum()) / tong * 100) if tong else 0}%")
        k[2].metric("Tổng lượt vắng", int(tk["Vắng"].sum()))
        k[3].metric("Tổng lượt muộn", int(tk["Muộn"].sum()))
        st.dataframe(tk, hide_index=True, **FULL, height=420, column_config={
            "Chuyên cần (%)": st.column_config.ProgressColumn("Chuyên cần", min_value=0, max_value=100, format="%d%%")})
        if so_ngay:
            st.markdown("##### Bảng điểm danh chi tiết")
            ky_hieu = {"Đi học": "✓", "Muộn": "M", "Vắng": "V"}
            bang = dd_t.assign(Ngay=dd_t["Ngày"].map(lambda d: d.strftime("%d/%m")),
                               KH=dd_t["Trạng thái"].map(ky_hieu)).pivot_table(
                index="STT", columns="Ngay", values="KH", aggfunc="last")
            bang = hoc_sinh[["STT", "Họ và tên"]].set_index("STT").join(bang).fillna("")
            st.dataframe(bang, **FULL, height=420)
            st.caption("✓ đi học · M muộn · V vắng")

    # ---------------- Thông tin học sinh ----------------
    with tab_hs:
        st.caption("Sửa trực tiếp trong bảng (ngày sinh, SĐT phụ huynh, nhóm, tài khoản…). Thêm học sinh mới ở dòng cuối – "
                   "tài khoản và mật khẩu được tự cấp khi lưu. Bấm **Lưu thông tin học sinh** để ghi lại.")
        sua = st.data_editor(hoc_sinh, hide_index=True, **FULL, num_rows="dynamic", key="sua_hs",
                             height=520, column_config={
                                 "STT": st.column_config.NumberColumn("STT", disabled=True, width="small"),
                                 "Họ và tên": st.column_config.TextColumn("Họ và tên", required=True, width="medium"),
                                 "Ngày sinh": st.column_config.TextColumn("Ngày sinh", help="Ví dụ 05/09/2014"),
                                 "SĐT phụ huynh": st.column_config.TextColumn("SĐT phụ huynh"),
                                 "Tên đăng nhập": st.column_config.TextColumn("Tên đăng nhập", help="Tên + số thứ tự, VD an01"),
                                 "Mật khẩu": st.column_config.TextColumn("Mật khẩu", help="2 chữ số"),
                             })
        if st.button("Lưu thông tin học sinh", type="primary"):
            sua = sua[sua["Họ và tên"].fillna("").astype(str).str.strip() != ""].copy()
            moi = sua["STT"].isna()
            if moi.any():
                bat_dau = int(pd.to_numeric(sua["STT"], errors="coerce").max() or 0) + 1
                sua.loc[moi, "STT"] = range(bat_dau, bat_dau + int(moi.sum()))
            sua["STT"] = sua["STT"].astype(int)
            for c in ["Ngày sinh", "SĐT phụ huynh", "Nhóm/Lớp", "Ghi chú", "Tên đăng nhập", "Mật khẩu"]:
                sua[c] = sua[c].map(dl._chu)
            sua = dl.cap_tai_khoan(sua)
            trung = sua["Tên đăng nhập"][sua["Tên đăng nhập"].duplicated()].tolist()
            if trung:
                st.error(f"Tên đăng nhập bị trùng: {', '.join(trung)}. Sửa lại cho khác nhau rồi lưu.")
            else:
                d = nap_du_lieu()
                d["HocSinh"] = sua[dl.COT["HocSinh"]]
                if luu(d, "Đã lưu thông tin học sinh", ["HocSinh"]):
                    st.rerun()
        with st.expander("🔑 Danh sách tài khoản để phát cho học sinh"):
            st.code("\n".join(f"{int(r['STT'])}. {r['Họ và tên']} – Tên đăng nhập: {r['Tên đăng nhập']} – Mật khẩu: {r['Mật khẩu']}"
                              for _, r in hoc_sinh.iterrows()), language=None)
            st.caption("Muốn đổi mật khẩu một em: sửa ô Mật khẩu trong bảng trên rồi bấm Lưu.")
            if st.button("Cấp lại mật khẩu ngẫu nhiên cho cả lớp"):
                st.session_state["xac_nhan_cap_lai"] = True
            if st.session_state.get("xac_nhan_cap_lai"):
                st.warning("Mật khẩu cũ của mọi em sẽ không dùng được nữa.")
                if st.button("Đồng ý cấp lại", type="primary"):
                    d = nap_du_lieu()
                    d["HocSinh"] = dl.cap_tai_khoan(d["HocSinh"], cap_lai_mat_khau=True)
                    st.session_state["xac_nhan_cap_lai"] = False
                    if luu(d, "Đã cấp lại mật khẩu cả lớp", ["HocSinh"]):
                        st.rerun()


# ============================================================================
# MÀN HÌNH 2 – NHẬP & THEO DÕI ĐIỂM
# ============================================================================
def bieu_do_diem(d: pd.DataFrame, tb):
    """Biểu đồ đường điểm qua các bài, có đường điểm trung bình."""
    d = d.copy()
    d["Nhãn"] = [f"{i}. {r['Ngày']:%d/%m} · {(r['Tên bài'] or r['Loại bài'])[:22]}" for i, (_, r) in enumerate(d.iterrows(), 1)]
    co_so = alt.Chart(d).encode(x=alt.X("Nhãn:N", sort=None, title=None, axis=alt.Axis(labelAngle=-30, labelLimit=180)))
    duong = co_so.mark_line(color="#C0392B", strokeWidth=2.5).encode(y=alt.Y("Điểm:Q", scale=alt.Scale(domain=[0, 10]), title="Điểm"))
    diem = co_so.mark_circle(size=110, opacity=1).encode(
        y="Điểm:Q", color=alt.Color("Loại bài:N", scale=alt.Scale(scheme="set1"), legend=alt.Legend(orient="top", title=None)),
        tooltip=["Nhãn", "Loại bài", alt.Tooltip("Điểm:Q", format=".2f"), "Nhận xét"])
    chu = co_so.mark_text(dy=-14, fontSize=12, color="#1B2430").encode(y="Điểm:Q", text=alt.Text("Điểm:Q", format=".3~f"))
    lop = [duong, diem, chu]
    if tb is not None:
        lop.append(alt.Chart(pd.DataFrame({"TB": [tb]})).mark_rule(strokeDash=[6, 4], color="#1F4E79").encode(y="TB:Q"))
    return alt.layer(*lop).properties(height=320, width="container")


def man_hinh_diem():
    st.title("Nhập & theo dõi điểm")
    tab1, tab2, tab3 = st.tabs(["👤 Theo từng học sinh", "📝 Nhập điểm cả lớp cho một bài", "📊 Bảng điểm cả lớp"])

    # ---------------- Theo từng học sinh ----------------
    with tab1:
        stt = st.selectbox("Chọn học sinh", DS_STT, format_func=lambda s: nhan_hs(data, s), key="hs_diem")
        with st.form("form_diem", clear_on_submit=True):
            c = st.columns([1.1, 1.1, 2, 0.9])
            ngay = c[0].date_input("Ngày kiểm tra", date.today(), format="DD/MM/YYYY")
            loai = c[1].selectbox("Loại bài", dl.LOAI_BAI)
            ten_bai = c[2].text_input("Tên bài", placeholder="VD: Định kỳ giữa kỳ I / Chuyên đề Phân số")
            diem = c[3].number_input("Điểm", min_value=0.0, max_value=10.0, step=0.25, value=None, placeholder="0–10")
            nx = st.text_input("Nhận xét bài này (không bắt buộc)")
            if st.form_submit_button("➕ Lưu điểm", type="primary"):
                if diem is None:
                    st.warning("Hãy nhập điểm.")
                else:
                    d = nap_du_lieu()
                    dl.them_diem(d, ngay, stt, loai, ten_bai.strip(), diem, nx)
                    if luu(d, f"Đã lưu điểm {dl.dinh_dang_diem(diem)} cho {dl.ten_hs(d, stt)}", ["Diem"]):
                        st.rerun()

        d_hs = dl.diem_cua_hs(data, stt)
        tt = dl.tom_tat_diem(d_hs)
        k = st.columns(5)
        k[0].metric("Điểm TB chung", dl.dinh_dang_diem(tt["tb"]), dl.xep_loai(tt["tb"]), delta_color="off")
        k[1].metric("TB Định kỳ", dl.dinh_dang_diem(tt["tb_loai"].get("Định kỳ")))
        k[2].metric("TB Chuyên đề", dl.dinh_dang_diem(tt["tb_loai"].get("Chuyên đề")))
        k[3].metric("Bài gần nhất", dl.dinh_dang_diem(tt["gan_nhat"]),
                    None if tt["thay_doi"] is None else f"{tt['thay_doi']:+.2f} so với bài trước".replace(".", ","))
        k[4].metric("Xu hướng", tt["xu_huong"], f"{tt['so_bai']} bài", delta_color="off")

        if d_hs.empty:
            st.info("Học sinh này chưa có điểm. Nhập bài đầu tiên ở khung phía trên.")
        else:
            st.altair_chart(bieu_do_diem(d_hs, tt["tb"]))
            with st.expander("Sửa / xoá điểm của học sinh này"):
                st.caption("Sửa trực tiếp trong bảng; chọn dòng rồi bấm thùng rác để xoá. Sau đó bấm Lưu thay đổi.")
                sua = st.data_editor(d_hs[["Ngày", "Loại bài", "Tên bài", "Điểm", "Nhận xét"]], hide_index=True,
                                     num_rows="dynamic", **FULL, key=f"sua_diem_{stt}", column_config={
                                         "Ngày": st.column_config.DateColumn("Ngày", format="DD/MM/YYYY", required=True),
                                         "Loại bài": st.column_config.SelectboxColumn("Loại bài", options=dl.LOAI_BAI, required=True),
                                         "Điểm": st.column_config.NumberColumn("Điểm", min_value=0, max_value=10, step=0.25, required=True)})
                if st.button("Lưu thay đổi điểm"):
                    d = nap_du_lieu()
                    khac = d["Diem"][d["Diem"]["STT"] != stt]
                    sua = sua.dropna(subset=["Ngày", "Điểm"]).copy()
                    sua["STT"] = stt
                    sua["Họ và tên"] = dl.ten_hs(d, stt)
                    sua["Ngày"] = sua["Ngày"].map(dl._ve_ngay)
                    sua["Tên bài"] = sua["Tên bài"].fillna("")
                    sua["Nhận xét"] = sua["Nhận xét"].fillna("")
                    d["Diem"] = pd.concat([khac, sua[dl.COT["Diem"]]], ignore_index=True).sort_values(["Ngày", "STT"])
                    if luu(d, "Đã cập nhật bảng điểm", ["Diem"]):
                        st.rerun()

    # ---------------- Nhập cả lớp ----------------
    with tab2:
        c = st.columns([1.1, 1.1, 2.2])
        ngay2 = c[0].date_input("Ngày kiểm tra", date.today(), format="DD/MM/YYYY", key="ngay_ca_lop")
        loai2 = c[1].selectbox("Loại bài", dl.LOAI_BAI, key="loai_ca_lop")
        ten2 = c[2].text_input("Tên bài", placeholder="VD: Kiểm tra chuyên đề Số nguyên", key="ten_ca_lop")
        dd = data["Diem"]
        cu = dd[(dd["Ngày"] == ngay2) & (dd["Tên bài"] == ten2.strip())].set_index("STT")
        bang = hoc_sinh[["STT", "Họ và tên"]].copy()
        bang["Điểm"] = bang["STT"].map(cu["Điểm"]).astype(float) if len(cu) else pd.Series([float("nan")] * len(bang), dtype=float)
        bang["Nhận xét"] = bang["STT"].map(cu["Nhận xét"]).fillna("").astype(str) if len(cu) else ""
        st.caption("Gõ điểm vào cột **Điểm** (bấm Enter để xuống em tiếp theo). Em nào để trống thì không lưu.")
        nhap = st.data_editor(bang, hide_index=True, **FULL, height=560,
                              key=f"ca_lop_{ngay2}_{ten2}", column_config={
                                  "STT": st.column_config.NumberColumn(disabled=True, width="small"),
                                  "Họ và tên": st.column_config.TextColumn(disabled=True),
                                  "Điểm": st.column_config.NumberColumn(min_value=0, max_value=10, step=0.25),
                                  "Nhận xét": st.column_config.TextColumn(width="large")})
        co_diem = nhap["Điểm"].notna()
        if co_diem.any():
            s = nhap.loc[co_diem, "Điểm"]
            k = st.columns(4)
            k[0].metric("Đã nhập", f"{int(co_diem.sum())}/{len(nhap)}")
            k[1].metric("Điểm TB bài", dl.dinh_dang_diem(s.mean()))
            k[2].metric("Cao nhất / Thấp nhất", f"{dl.dinh_dang_diem(s.max())} / {dl.dinh_dang_diem(s.min())}")
            k[3].metric("Dưới 5 điểm", int((s < 5).sum()))
        if st.button("💾 Lưu điểm cả lớp", type="primary", disabled=not ten2.strip()):
            d = nap_du_lieu()
            for _, r in nhap.iterrows():
                if pd.notna(r["Điểm"]):
                    dl.them_diem(d, ngay2, int(r["STT"]), loai2, ten2.strip(), float(r["Điểm"]), r["Nhận xét"] or "")
            if luu(d, f"Đã lưu điểm bài '{ten2.strip()}' cho {int(co_diem.sum())} học sinh", ["Diem"]):
                st.rerun()
        if not ten2.strip():
            st.caption("Nhập **Tên bài** để bật nút Lưu.")

    # ---------------- Bảng điểm cả lớp ----------------
    with tab3:
        dong = []
        for stt in DS_STT:
            t = dl.tom_tat_diem(dl.diem_cua_hs(data, stt))
            dong.append({"STT": stt, "Họ và tên": dl.ten_hs(data, stt), "Số bài": t["so_bai"], "TB chung": t["tb"],
                         "TB Định kỳ": t["tb_loai"].get("Định kỳ"), "TB Chuyên đề": t["tb_loai"].get("Chuyên đề"),
                         "Bài gần nhất": t["gan_nhat"], "Xu hướng": t["xu_huong"], "Xếp loại": dl.xep_loai(t["tb"])})
        bang = pd.DataFrame(dong)
        st.dataframe(bang, hide_index=True, **FULL, height=600, column_config={
            "TB chung": st.column_config.ProgressColumn("TB chung", min_value=0, max_value=10, format="%.2f")})


# ============================================================================
# MÀN HÌNH 3 – HỌC PHÍ & THU – CHI
# ============================================================================
def man_hinh_hoc_phi():
    st.title("Học phí & Thu – Chi")
    thang = chon_thang("Tháng", "thang_hp")
    bang = dl.bang_hoc_phi(data, thang)
    chi = dl.chi_trong_thang(data, thang)
    da_thu = float(bang["Đã đóng"].fillna(0).sum())
    con_no = float(bang["Còn nợ"].sum())
    tong_chi = float(chi["Số tiền"].fillna(0).sum())
    k = st.columns(5)
    k[0].metric("Đã đóng đủ", f"{(bang['Trạng thái'] == 'Đã đóng đủ').sum()}/{len(bang)}")
    k[1].metric("Đã thu", dl.dinh_dang_tien(da_thu))
    k[2].metric("Còn phải thu", dl.dinh_dang_tien(con_no))
    k[3].metric("Đã chi", dl.dinh_dang_tien(tong_chi))
    k[4].metric("Còn lại (Thu – Chi)", dl.dinh_dang_tien(da_thu - tong_chi))

    tab_hp, tab_chi, tab_tong = st.tabs(["💰 Học phí tháng", "🧾 Khoản chi", "📈 Tổng hợp các tháng"])
    with tab_hp:
        loc = st.radio("Hiển thị", ["Tất cả", "Chưa đóng", "Còn nợ", "Đã đóng đủ"], horizontal=True, key="loc_hp")
        xem = bang if loc == "Tất cả" else bang[bang["Trạng thái"] == loc]
        xem = xem.assign(**{"Đóng đủ": xem["Trạng thái"].isin(["Đã đóng đủ", "Miễn phí"])})
        st.caption("Tích **Đóng đủ** để ghi nhận đã thu đủ (tự điền số tiền và ngày hôm nay), hoặc gõ số tiền đã đóng nếu "
                   "đóng một phần. Bấm **Lưu học phí** để ghi vào Excel.")
        sua = st.data_editor(
            xem[["STT", "Họ và tên", "Đóng đủ", "Phải đóng", "Đã đóng", "Ngày đóng", "Trạng thái", "Ghi chú"]],
            hide_index=True, **FULL, height=520, key=f"hp_{thang}_{loc}",
            column_config={
                "STT": st.column_config.NumberColumn(disabled=True, width="small"),
                "Họ và tên": st.column_config.TextColumn(disabled=True),
                "Đóng đủ": st.column_config.CheckboxColumn("Đóng đủ", width="small"),
                "Phải đóng": st.column_config.NumberColumn(min_value=0, step=10000, format="%d"),
                "Đã đóng": st.column_config.NumberColumn(min_value=0, step=10000, format="%d"),
                "Ngày đóng": st.column_config.DateColumn(format="DD/MM/YYYY"),
                "Trạng thái": st.column_config.TextColumn(disabled=True),
            })
        if st.button("💾 Lưu học phí", type="primary"):
            sua = sua.copy()
            tich = sua["Đóng đủ"].fillna(False).astype(bool)
            sua.loc[tich & (sua["Đã đóng"].fillna(0) < sua["Phải đóng"].fillna(0)), "Đã đóng"] = sua["Phải đóng"]
            sua.loc[tich & sua["Ngày đóng"].isna(), "Ngày đóng"] = date.today()
            # Ghép phần vừa sửa vào bảng đầy đủ của tháng (theo STT)
            day_du = bang.set_index("STT")
            for c in ["Phải đóng", "Đã đóng", "Ngày đóng", "Ghi chú"]:
                day_du[c] = day_du[c].astype(object)
                day_du.loc[sua["STT"].values, c] = sua[c].values
            d = nap_du_lieu()
            dl.luu_hoc_phi(d, thang, day_du.reset_index())
            if luu(d, f"Đã lưu học phí {dl.nhan_thang(thang)}", ["HocPhi"]):
                st.rerun()
        no = bang[bang["Còn nợ"] > 0]
        if len(no):
            with st.expander(f"📋 Danh sách nhắc học phí ({len(no)} em) – sao chép để gửi"):
                st.code("\n".join([f"Nhắc học phí {dl.nhan_thang(thang).lower()}:"] +
                                  [f"{int(r['STT'])}. {r['Họ và tên']} – còn {dl.dinh_dang_tien(r['Còn nợ'])}" for _, r in no.iterrows()]),
                        language=None)

    with tab_chi:
        st.caption("Ghi các khoản chi trong tháng (photo đề, thuê phòng, quà thưởng…). Thêm dòng mới ở cuối bảng.")
        chi_bang = chi.copy() if len(chi) else pd.DataFrame({"Ngày": pd.Series(dtype="datetime64[ns]"), "Nội dung": pd.Series(dtype="object"),
                                                              "Số tiền": pd.Series(dtype="float"), "Ghi chú": pd.Series(dtype="object")})
        if len(chi_bang):
            chi_bang["Ngày"] = pd.to_datetime(chi_bang["Ngày"])
            chi_bang["Số tiền"] = chi_bang["Số tiền"].astype(float)
        sua_chi = st.data_editor(chi_bang,
                                 hide_index=True, num_rows="dynamic", **FULL, key=f"chi_{thang}",
                                 column_config={"Ngày": st.column_config.DateColumn(format="DD/MM/YYYY", default=date.today()),
                                                "Số tiền": st.column_config.NumberColumn(min_value=0, step=10000, format="%d")})
        if st.button("💾 Lưu khoản chi"):
            d = nap_du_lieu()
            khac = d["ChiPhi"][d["ChiPhi"]["Ngày"].map(lambda x: x.strftime("%Y-%m") if x else "") != thang]
            moi = sua_chi.dropna(subset=["Số tiền"]).copy()
            moi["Ngày"] = moi["Ngày"].map(dl._ve_ngay)
            moi["Ngày"] = moi["Ngày"].map(lambda x: x or date.today())
            moi["Nội dung"] = moi["Nội dung"].fillna("")
            moi["Ghi chú"] = moi["Ghi chú"].fillna("")
            d["ChiPhi"] = pd.concat([khac, moi[dl.COT["ChiPhi"]]], ignore_index=True).sort_values("Ngày")
            if luu(d, "Đã lưu khoản chi", ["ChiPhi"]):
                st.rerun()

    with tab_tong:
        cac_thang = sorted(set(data["HocPhi"]["Tháng"]) | {x.strftime("%Y-%m") for x in data["ChiPhi"]["Ngày"].dropna()})
        if not cac_thang:
            st.info("Chưa có dữ liệu học phí hoặc khoản chi nào được lưu.")
        else:
            dong = []
            for t in cac_thang:
                hp = data["HocPhi"][data["HocPhi"]["Tháng"] == t]
                thu = float(hp["Đã đóng"].fillna(0).sum())
                c = float(dl.chi_trong_thang(data, t)["Số tiền"].fillna(0).sum())
                dong.append({"Tháng": dl.nhan_thang(t), "Thu": thu, "Chi": c, "Còn lại": thu - c})
            tong = pd.DataFrame(dong)
            st.dataframe(tong.style.format({"Thu": dl.dinh_dang_tien, "Chi": dl.dinh_dang_tien, "Còn lại": dl.dinh_dang_tien}),
                         hide_index=True, **FULL)
            st.bar_chart(tong.set_index("Tháng")[["Thu", "Chi"]], color=["#1F4E79", "#C0392B"], stack=False)


# ============================================================================
# MÀN HÌNH 4 – BÁO CÁO PHỤ HUYNH
# ============================================================================
def goi_y_nhan_xet(bc) -> str:
    """Gợi ý nhận xét tự động từ số liệu (bạn sửa lại cho phù hợp)."""
    cc, tt = bc["chuyen_can"], bc["tom_tat"]
    ten = bc["ten"].split()[-1]
    cau = []
    if cc["Số buổi"]:
        if cc["Vắng"] == 0 and cc["Muộn"] == 0:
            cau.append(f"Tháng này {ten} đi học đầy đủ, đúng giờ.")
        elif cc["Vắng"] == 0:
            cau.append(f"{ten} đi học đầy đủ nhưng còn đi muộn {int(cc['Muộn'])} buổi.")
        else:
            cau.append(f"{ten} vắng {int(cc['Vắng'])} buổi, gia đình nhắc con đi học đều hơn.")
    if tt["tb"] is not None:
        cau.append(f"Điểm trung bình các bài kiểm tra đạt {dl.dinh_dang_diem(tt['tb'])} ({dl.xep_loai(tt['tb']).lower()}).")
        if tt["xu_huong"] == "Tiến bộ":
            cau.append("Kết quả có tiến bộ rõ so với đầu tháng, con cần giữ vững phong độ.")
        elif tt["xu_huong"] == "Đi xuống":
            cau.append("Điểm các bài gần đây giảm, con cần ôn lại kiến thức và làm bài tập về nhà đầy đủ.")
        if tt["tb"] < 6.5:
            cau.append("Con cần luyện thêm các dạng bài cơ bản và trình bày lời giải cẩn thận hơn.")
    return " ".join(cau)


def man_hinh_bao_cao():
    st.title("Báo cáo gửi phụ huynh")
    c1, c2, c3 = st.columns([2, 1.2, 1])
    stt = c1.selectbox("Học sinh", DS_STT, format_func=lambda s: nhan_hs(data, s), key="hs_bc")
    with c2:
        thang = chon_thang("Tháng báo cáo", "thang_bc")
    with c3:
        st.write("")
        kem_hp = st.checkbox("Kèm học phí", value=True)
    bc = dl.noi_dung_bao_cao(data, stt, thang, kem_hp)

    # Nhận xét cuối tháng
    khoa_nx = f"nx_{stt}_{thang}"
    if khoa_nx not in st.session_state:
        st.session_state[khoa_nx] = bc["nhan_xet"]

    def _goi_y():
        st.session_state[khoa_nx] = goi_y_nhan_xet(bc)

    st.text_area("Nhận xét của giáo viên (tháng này)", key=khoa_nx, height=110,
                 placeholder="Ví dụ: Con nắm chắc phần phân số, cần cẩn thận hơn khi tính toán…")
    b1, b2, _ = st.columns([1.3, 1.6, 4])
    b2.button("✨ Gợi ý nhận xét từ số liệu", on_click=_goi_y, **FULL)
    if b1.button("💾 Lưu nhận xét", type="primary", **FULL):
        d = nap_du_lieu()
        dl.ghi_nhan_xet(d, thang, stt, st.session_state[khoa_nx])
        if luu(d, "Đã lưu nhận xét", ["NhanXet"]):
            st.rerun()
    bc["nhan_xet"] = st.session_state[khoa_nx]

    st.divider()
    trai, phai = st.columns([1, 1], gap="large")
    with trai:
        st.markdown("##### Tin nhắn (dán vào Zalo / Messenger)")
        van_ban = dl.bao_cao_van_ban(bc)
        st.code(van_ban, language=None)
        st.caption("Bấm biểu tượng 📋 ở góc khung để sao chép.")
        st.download_button("⬇️ Tải file chữ (.txt)", van_ban.encode("utf-8"),
                           file_name=f"BaoCao_{bc['ten']}_{thang}.txt", mime="text/plain")
    with phai:
        st.markdown("##### Phiếu báo cáo (ảnh)")
        anh = dl.bao_cao_anh(bc)
        st.image(anh, width=560)
        st.download_button("⬇️ Tải ảnh phiếu báo cáo (.png)", anh, file_name=f"BaoCao_{bc['ten']}_{thang}.png",
                           mime="image/png", type="primary")

    st.divider()
    with st.expander("📦 Xuất phiếu báo cáo cho cả lớp (1 file .zip)"):
        st.caption("Tạo ảnh báo cáo cho tất cả học sinh trong tháng đã chọn. Mất khoảng 10–20 giây.")
        if st.button("Tạo file zip"):
            buf = io.BytesIO()
            with st.spinner("Đang tạo phiếu báo cáo…"), zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
                for s in DS_STT:
                    b = dl.noi_dung_bao_cao(data, s, thang, kem_hp)
                    z.writestr(f"{s:02d}_{b['ten']}_{thang}.png", dl.bao_cao_anh(b))
                    z.writestr(f"{s:02d}_{b['ten']}_{thang}.txt", dl.bao_cao_van_ban(b))
            st.download_button("⬇️ Tải file zip", buf.getvalue(), file_name=f"BaoCao_CaLop_{thang}.zip",
                               mime="application/zip", type="primary")


# ============================================================================
# CÀI ĐẶT
# ============================================================================
def man_hinh_cai_dat():
    st.title("Cài đặt")
    st.caption("Các thông tin này hiện trên báo cáo gửi phụ huynh và dùng làm học phí mặc định.")
    sua = st.data_editor(data["CaiDat"], hide_index=True, **FULL, key="sua_cd",
                         column_config={"Mục": st.column_config.TextColumn(disabled=True),
                                        "Giá trị": st.column_config.TextColumn(width="large")})
    if st.button("💾 Lưu cài đặt", type="primary"):
        d = nap_du_lieu()
        d["CaiDat"] = sua.fillna("")
        if luu(d, "Đã lưu cài đặt", ["CaiDat"]):
            st.rerun()
    st.divider()
    if ONLINE:
        st.markdown(f"**Dữ liệu lưu tại Google Sheets:** [mở bảng tính]({BI_MAT['gsheets']['url']})")
        st.caption("Google Sheets tự lưu lịch sử phiên bản: Tệp → Nhật ký phiên bản – dùng khi cần khôi phục.")
    else:
        st.markdown(f"**File dữ liệu:** `{dl.FILE_DU_LIEU}`")
        st.markdown(f"**Thư mục sao lưu:** `{dl.THU_MUC_SAO_LUU}`")
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        for ten in dl.COT:
            data[ten].to_excel(w, sheet_name=ten, index=False)
    st.download_button("⬇️ Tải bản sao dữ liệu (Excel)", buf.getvalue(), file_name=f"SaoLuu_LopToan_{date.today():%Y-%m-%d}.xlsx")


# ============================================================================
# ĐĂNG NHẬP (chỉ khi chạy online)
# ============================================================================
def man_hinh_dang_nhap():
    _, giua, _ = st.columns([1, 1.3, 1])
    with giua:
        st.markdown(f"<div style='text-align:center;margin-top:2rem'><div class='nho'>{cd.get('Tên trung tâm', '')}</div>"
                    f"<h1 style='margin:.2rem 0 1rem'>📐 Lớp {cd.get('Môn', 'Toán')} – Đăng nhập</h1></div>", unsafe_allow_html=True)
        with st.form("dang_nhap"):
            ten = st.text_input("Tên đăng nhập", placeholder="VD: an01 (tên + số thứ tự)")
            mk = st.text_input("Mật khẩu", type="password", placeholder="2 chữ số")
            bam = st.form_submit_button("Đăng nhập", type="primary")
        if bam:
            sai, luc = st.session_state.get("_sai", 0), st.session_state.get("_sai_luc", 0.0)
            if sai >= 5 and time.time() - luc < 60:
                st.error("Bạn nhập sai quá 5 lần. Chờ 1 phút rồi thử lại.")
            else:
                kq = dl.kiem_tra_dang_nhap(data, ten, mk, TK_GV, MK_GV)
                if kq:
                    st.session_state["nguoi_dung"] = kq
                    st.session_state["_sai"] = 0
                    st.rerun()
                st.session_state["_sai"] = sai + 1 if time.time() - luc < 60 else 1
                st.session_state["_sai_luc"] = time.time()
                st.error("Sai tên đăng nhập hoặc mật khẩu.")
        st.caption("Tên đăng nhập là **tên của em + số thứ tự** trong danh sách lớp (ví dụ An, STT 1 → **an01**). "
                   "Quên mật khẩu thì hỏi thầy/cô.")


def dang_xuat():
    for k in ("nguoi_dung",):
        st.session_state.pop(k, None)


# ============================================================================
# TRANG CỦA HỌC SINH (chỉ xem kết quả của chính mình)
# ============================================================================
def man_hinh_hoc_sinh(stt: int):
    hs = hoc_sinh[hoc_sinh["STT"] == stt]
    if hs.empty:
        st.error("Tài khoản này không còn trong danh sách lớp. Hãy hỏi thầy/cô.")
        st.button("Đăng xuất", on_click=dang_xuat)
        return
    hs = hs.iloc[0]
    trai, phai = st.columns([4, 1])
    trai.markdown(f"<div class='nho'>{cd.get('Tên trung tâm', '')} · {hs['Nhóm/Lớp']} · STT {stt}</div>", unsafe_allow_html=True)
    trai.title(hs["Họ và tên"])
    phai.write("")
    phai.button("Đăng xuất", on_click=dang_xuat, **FULL)

    thang = chon_thang("Tháng", "thang_hs")
    cc = dl.thong_ke_chuyen_can(data, thang)
    cc = cc[cc["STT"] == stt].iloc[0]
    d_thang = dl.diem_cua_hs(data, stt, thang)
    tt = dl.tom_tat_diem(d_thang)
    d_tat_ca = dl.diem_cua_hs(data, stt)
    tt_all = dl.tom_tat_diem(d_tat_ca)

    k = st.columns(4)
    k[0].metric("Đi học", f"{int(cc['Đi học']) + int(cc['Muộn'])}/{int(cc['Số buổi'])} buổi",
                f"Vắng {int(cc['Vắng'])} · Muộn {int(cc['Muộn'])}", delta_color="off")
    k[1].metric("Điểm TB tháng", dl.dinh_dang_diem(tt["tb"]), dl.xep_loai(tt["tb"]) if tt["tb"] is not None else None, delta_color="off")
    k[2].metric("Điểm TB từ đầu", dl.dinh_dang_diem(tt_all["tb"]), f"{tt_all['so_bai']} bài", delta_color="off")
    k[3].metric("Xu hướng", tt_all["xu_huong"])

    nx = dl.lay_nhan_xet(data, thang, stt)
    st.markdown("##### Nhận xét của thầy/cô")
    st.info(nx or "Thầy/cô chưa có nhận xét cho tháng này.")

    st.markdown("##### Điểm các bài kiểm tra (từ đầu khoá)")
    if d_tat_ca.empty:
        st.caption("Chưa có bài kiểm tra nào.")
    else:
        st.altair_chart(bieu_do_diem(d_tat_ca, tt_all["tb"]))
    if not d_thang.empty:
        st.markdown(f"##### Bài kiểm tra {dl.nhan_thang(thang).lower()}")
        bang = d_thang[["Ngày", "Tên bài", "Loại bài", "Điểm", "Nhận xét"]].copy()
        bang["Ngày"] = bang["Ngày"].map(lambda x: x.strftime("%d/%m/%Y") if x else "")
        st.dataframe(bang, hide_index=True, **FULL)

    dd = data["DiemDanh"]
    dd = dd[(dd["STT"] == stt) & (dd["Ngày"].map(lambda x: x.strftime("%Y-%m") if x else "") == thang)]
    if len(dd):
        st.markdown("##### Điểm danh trong tháng")
        st.dataframe(pd.DataFrame({"Ngày": dd["Ngày"].map(lambda x: x.strftime("%d/%m/%Y")), "Trạng thái": dd["Trạng thái"]}),
                     hide_index=True, **FULL)

    xem_hp = str(cd.get("Học sinh xem được học phí", "Có")).strip().lower() in ("có", "co", "yes", "1", "x")
    bc = dl.noi_dung_bao_cao(data, stt, thang, xem_hp)
    if xem_hp and bc["hoc_phi"]:
        hp = bc["hoc_phi"]
        st.markdown("##### Học phí")
        (st.success if hp["Còn nợ"] <= 0 else st.warning)(
            f"{dl.nhan_thang(thang)}: {hp['Trạng thái']}" + (f" – còn {dl.dinh_dang_tien(hp['Còn nợ'])}" if hp["Còn nợ"] > 0 else ""))
    st.download_button("⬇️ Tải phiếu báo cáo tháng (ảnh)", dl.bao_cao_anh(bc),
                       file_name=f"BaoCao_{bc['ten']}_{thang}.png", mime="image/png")


# ============================================================================
# CHẠY ỨNG DỤNG
# ============================================================================
if ONLINE:
    nguoi_dung = st.session_state.get("nguoi_dung")
    if not nguoi_dung:
        man_hinh_dang_nhap()
        st.stop()
else:
    nguoi_dung = ("gv", None)            # chạy trên máy: vào thẳng trang giáo viên

if nguoi_dung[0] == "hs":
    man_hinh_hoc_sinh(nguoi_dung[1])
    st.stop()

if hoc_sinh.empty:                      # Google Sheet mới tạo, chưa có học sinh
    man_hinh_khoi_tao()
    st.stop()

with st.sidebar:
    st.markdown(f"### 📐 {cd.get('Tên trung tâm', 'Lớp Toán')}")
    st.caption(f"{len(hoc_sinh)} học sinh · {cd.get('Giáo viên', '')}")
    trang = st.radio("Màn hình", ["1. Tổng quan & Điểm danh", "2. Nhập & theo dõi điểm",
                                  "3. Học phí & Thu – Chi", "4. Báo cáo phụ huynh", "⚙️ Cài đặt"],
                     label_visibility="collapsed")
    st.divider()
    if ONLINE:
        st.caption("Dữ liệu lưu trên **Google Sheets**.")
        st.button("Đăng xuất", on_click=dang_xuat)
    else:
        st.caption(f"Dữ liệu: `{dl.FILE_DU_LIEU.name}`  \nSao lưu tự động mỗi ngày trong thư mục `sao_luu`.")
        st.caption("Nhớ **đóng file Excel** trước khi lưu trong ứng dụng.")

if trang.startswith("1"):
    man_hinh_diem_danh()
elif trang.startswith("2"):
    man_hinh_diem()
elif trang.startswith("3"):
    man_hinh_hoc_phi()
elif trang.startswith("4"):
    man_hinh_bao_cao()
else:
    man_hinh_cai_dat()
